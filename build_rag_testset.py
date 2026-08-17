from pathlib import Path
import json
import math
import shutil

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "rag_multimodal_testset"
ASSETS = OUT / "assets"
DOCS = OUT / "documents"


def font(size, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size=size)
    return ImageFont.load_default()


def draw_center(draw, xy, text, f, fill):
    box = draw.textbbox((0, 0), text, font=f)
    w, h = box[2] - box[0], box[3] - box[1]
    draw.text((xy[0] - w / 2, xy[1] - h / 2), text, font=f, fill=fill)


def make_scene(path):
    img = Image.new("RGB", (1400, 800), "#eef4f7")
    d = ImageDraw.Draw(img)
    # Warehouse shell and floor guides.
    d.rectangle((0, 0, 1400, 510), fill="#dfe9ee")
    d.rectangle((0, 510, 1400, 800), fill="#cbd6dc")
    for x in range(0, 1401, 140):
        d.line((700, 510, x, 800), fill="#aebdc5", width=3)
    d.line((0, 510, 1400, 510), fill="#73858f", width=5)
    # Charging bay and green leaf sign.
    d.rectangle((70, 125, 335, 500), fill="#344955", outline="#1e3038", width=6)
    d.rectangle((115, 50, 290, 180), fill="#35a853", outline="white", width=6)
    d.ellipse((160, 78, 250, 140), fill="white")
    d.line((175, 140, 236, 88), fill="#35a853", width=8)
    # Exactly three blue robots.
    robot_x = [400, 650, 900]
    for i, x in enumerate(robot_x, 1):
        d.rounded_rectangle((x, 430, x + 180, 620), radius=34, fill="#1976d2", outline="#0c3d68", width=6)
        d.rectangle((x + 35, 470, x + 145, 535), fill="#bfe5ff", outline="#083b66", width=4)
        d.ellipse((x + 20, 595, x + 65, 640), fill="#263238")
        d.ellipse((x + 115, 595, x + 160, 640), fill="#263238")
        d.ellipse((x + 78, 555, x + 102, 579), fill="#7cff8b")
        draw_center(d, (x + 90, 505), f"R-{i}", font(28, True), "#083b66")
    # Exactly two orange crates.
    for y in (410, 560):
        d.rectangle((1160, y, 1360, y + 130), fill="#f57c00", outline="#8a4200", width=7)
        d.line((1160, y, 1360, y + 130), fill="#ffb45b", width=5)
        d.line((1360, y, 1160, y + 130), fill="#ffb45b", width=5)
    # One worker, yellow helmet, tablet.
    d.ellipse((315, 260, 400, 345), fill="#f2bd91", outline="#6d4c41", width=4)
    d.pieslice((305, 235, 410, 330), 180, 360, fill="#ffd600", outline="#7d6700", width=4)
    d.polygon([(337, 340), (385, 340), (420, 485), (305, 485)], fill="#ffca28", outline="#6d5600")
    d.line((325, 370, 275, 455), fill="#6d4c41", width=18)
    d.rectangle((220, 420, 300, 520), fill="#263238", outline="#90caf9", width=5)
    d.line((392, 375, 445, 440), fill="#6d4c41", width=18)
    # Minimal title is deliberately generic; object facts must come from vision.
    draw_center(d, (700, 65), "智能仓库巡检现场", font(42, True), "#17324d")
    img.save(path, quality=95)


def make_bar_chart(path, title, labels, values, color="#3f7fc4"):
    w, h = 1200, 700
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f, label_f, value_f = font(38, True), font(27), font(27, True)
    draw_center(d, (w / 2, 55), title, title_f, "#17324d")
    left, top, right, bottom = 120, 130, 1110, 590
    d.line((left, top, left, bottom), fill="#36454f", width=4)
    d.line((left, bottom, right, bottom), fill="#36454f", width=4)
    maxv = math.ceil(max(values) / 500) * 500
    for tick in range(0, maxv + 1, 500):
        y = bottom - tick / maxv * (bottom - top)
        d.line((left, y, right, y), fill="#dbe3e8", width=2)
        d.text((30, y - 15), str(tick), font=label_f, fill="#44545c")
    gap = (right - left) / len(values)
    bw = gap * .55
    for i, (lab, val) in enumerate(zip(labels, values)):
        cx = left + gap * (i + .5)
        y = bottom - val / maxv * (bottom - top)
        d.rounded_rectangle((cx - bw / 2, y, cx + bw / 2, bottom), radius=10, fill=color)
        draw_center(d, (cx, y - 25), str(val), value_f, "#17324d")
        draw_center(d, (cx, bottom + 40), lab, label_f, "#17324d")
    img.save(path, quality=95)


def set_cell_text(cell, text, bold=False, size=13, color=None):
    cell.text = str(text)
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = PPt(size)
            run.font.bold = bold
            if color:
                run.font.color.rgb = color


def make_docx(scene, chart):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    props = doc.core_properties
    props.title = "星云智能仓库 2025 年运营报告"
    props.subject = "多模态 RAG 基准文档"
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"; styles["Normal"].font.size = Pt(10.5)
    for s in ["Title", "Heading 1", "Heading 2"]:
        styles[s].font.name = "Microsoft YaHei"
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("星云智能仓库 2025 年运营报告"); r.font.color.rgb = RGBColor(23, 50, 77)
    p = doc.add_paragraph("文档编号：NEBULA-DOCX-2025-01｜版本：1.0")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("1. 项目概况", level=1)
    doc.add_paragraph("“苍穹”试点位于苏州昆山，于 2025 年 3 月 18 日正式上线。项目负责人是林澈，年度目标是将单位订单能耗降低 12%，并把关键设备综合可用率维持在 98% 以上。")
    doc.add_heading("2. 巡检现场", level=1)
    doc.add_picture(str(scene), width=Inches(6.8))
    cap = doc.add_paragraph("图 1  智能仓库巡检现场（答案需基于图像对象识别）")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("现场图用于验证模型是否能从图片中识别对象数量、颜色、人员动作及标志，而不是只读取附近正文。")
    doc.add_heading("3. 关键设备清单", level=1)
    rows = [
        ["设备编号", "设备类型", "数量（台）", "可用率", "负责人"],
        ["R-01", "搬运机器人", "12", "99.2%", "林澈"],
        ["V-02", "视觉检测站", "4", "98.5%", "周宁"],
        ["C-03", "充电桩", "6", "97.8%", "韩梅"],
    ]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    for i, v in enumerate(rows[0]): table.rows[0].cells[i].text = v
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    doc.add_page_break()
    doc.add_heading("4. 季度订单趋势", level=1)
    doc.add_picture(str(chart), width=Inches(6.9))
    cap = doc.add_paragraph("图 2  2025 年季度完成订单量（单）")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("图表显示全年合计 5,850 单；第四季度达到峰值。复盘会决定将第四季度的排班策略用于下一年度。")
    doc.add_heading("5. 结论", level=1)
    doc.add_paragraph("试点进入稳定运行阶段。下一次正式复盘日期为 2026 年 1 月 12 日，会议室为海棠厅。")
    path = DOCS / "01_星云智能仓库运营报告.docx"
    doc.save(path)
    return path


def add_ppt_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(PInches(.6), PInches(.28), PInches(12.1), PInches(.75))
    p = tb.text_frame.paragraphs[0]; p.text = title
    p.font.name = "Microsoft YaHei"; p.font.size = PPt(28); p.font.bold = True; p.font.color.rgb = PRGBColor(23, 50, 77)
    if subtitle:
        sb = slide.shapes.add_textbox(PInches(.65), PInches(1.02), PInches(11.7), PInches(.45))
        p = sb.text_frame.paragraphs[0]; p.text = subtitle; p.font.name = "Microsoft YaHei"; p.font.size = PPt(13); p.font.color.rgb = PRGBColor(80, 95, 105)


def make_pptx(scene):
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "北辰物流中心扩容方案", "多模态 RAG 测试演示｜文档编号 NEBULA-PPTX-2025-02")
    slide.shapes.add_picture(str(scene), PInches(.7), PInches(1.65), width=PInches(6.4))
    box = slide.shapes.add_textbox(PInches(7.45), PInches(1.65), PInches(5.1), PInches(4.5))
    tf = box.text_frame; tf.word_wrap = True
    items = ["扩容代号：极光", "计划启用：2025 年 9 月 8 日", "新增作业区：B3 区", "安全负责人：陈曦", "高温预警阈值：38°C"]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = item; p.level = 0
        p.font.name = "Microsoft YaHei"; p.font.size = PPt(21); p.space_after = PPt(14)

    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "设备上线计划", "表格中的日期和责任人用于结构化检索测试")
    data = [
        ["阶段", "日期", "设备数", "责任人", "状态"],
        ["联调", "2025-08-12", "8", "顾言", "完成"],
        ["试运行", "2025-08-25", "16", "陈曦", "进行中"],
        ["正式启用", "2025-09-08", "24", "许舟", "待启动"],
    ]
    shape = slide.shapes.add_table(4, 5, PInches(.8), PInches(1.65), PInches(11.8), PInches(3.9))
    table = shape.table
    widths = [2.2, 2.5, 2.0, 2.0, 2.4]
    for i, width in enumerate(widths): table.columns[i].width = PInches(width)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c); set_cell_text(cell, val, bold=(r == 0), size=16)
            cell.fill.solid(); cell.fill.fore_color.rgb = PRGBColor(48, 92, 130) if r == 0 else PRGBColor(236, 243, 247)
            if r == 0:
                for p in cell.text_frame.paragraphs:
                    for run in p.runs: run.font.color.rgb = PRGBColor(255, 255, 255)

    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "每小时处理能力", "原生柱状图：单位为箱/小时")
    cd = ChartData(); cd.categories = ["A线", "B线", "C线"]; cd.add_series("处理能力", (420, 510, 465))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PInches(1.0), PInches(1.45), PInches(8.1), PInches(5.3), cd).chart
    chart.has_legend = False; chart.has_title = True; chart.chart_title.text_frame.text = "三条分拣线处理能力"
    chart.value_axis.maximum_scale = 600; chart.value_axis.minimum_scale = 0; chart.value_axis.major_unit = 100
    chart.plots[0].has_data_labels = True; chart.plots[0].data_labels.show_value = True
    box = slide.shapes.add_textbox(PInches(9.45), PInches(1.8), PInches(3.1), PInches(3.5))
    p = box.text_frame.paragraphs[0]; p.text = "结论\nB 线能力最高。\n\n异常联系人\n陈曦，分机 6027"
    p.font.name = "Microsoft YaHei"; p.font.size = PPt(20)
    path = DOCS / "02_北辰物流中心扩容方案.pptx"; prs.save(path); return path


def style_xlsx_title(ws, cell_range, title):
    ws.merge_cells(cell_range); cell = ws[cell_range.split(":")[0]]; cell.value = title
    cell.font = Font(name="Microsoft YaHei", size=20, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17324D"); cell.alignment = Alignment(horizontal="center", vertical="center")


def make_xlsx(scene):
    wb = Workbook(); ws = wb.active; ws.title = "项目概览"
    style_xlsx_title(ws, "A1:F2", "云帆仓配中心 2025 能效台账")
    info = [("项目代号", "青岚"), ("园区位置", "无锡新吴区"), ("能效负责人", "唐薇"), ("基准电价", "0.82 元/千瓦时"), ("年度节电目标", "18%")]
    for r, (k, v) in enumerate(info, 4):
        ws.cell(r, 1, k).font = Font(name="Microsoft YaHei", bold=True); ws.cell(r, 2, v).font = Font(name="Microsoft YaHei")
    ws["A10"] = "说明"; ws["A10"].font = Font(name="Microsoft YaHei", bold=True)
    ws["B10"] = "夜间低谷时段为 22:00–06:00；冷链机组优先在低谷时段预冷。"; ws["B10"].font = Font(name="Microsoft YaHei")
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 55

    data = wb.create_sheet("月度能耗")
    rows = [["月份", "订单量", "耗电量(kWh)", "单位能耗(kWh/单)", "峰值需量(kW)"],
            ["1月", 1000, 8200, 8.20, 310], ["2月", 1200, 9000, 7.50, 325],
            ["3月", 1500, 10200, 6.80, 340], ["4月", 1800, 11160, 6.20, 355]]
    for row in rows: data.append(row)
    for cell in data[1]:
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="2F6690"); cell.alignment = Alignment(horizontal="center")
    for row in data.iter_rows(min_row=2):
        for cell in row: cell.font = Font(name="Microsoft YaHei"); cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDE", [12, 14, 18, 22, 18]): data.column_dimensions[col].width = width
    chart = BarChart(); chart.type = "col"; chart.title = "月度单位能耗趋势"; chart.y_axis.title = "kWh/单"; chart.x_axis.title = "月份"
    chart.add_data(Reference(data, min_col=4, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=5)); chart.height = 8; chart.width = 14
    data.add_chart(chart, "G2")

    imgs = wb.create_sheet("现场图片")
    style_xlsx_title(imgs, "A1:H2", "仓库现场图像识别页")
    imgs["A4"] = "请从图片本身识别机器人、货箱、人员动作和绿色标志。"; imgs["A4"].font = Font(name="Microsoft YaHei", italic=True)
    xi = XLImage(str(scene)); xi.width = 840; xi.height = 480; imgs.add_image(xi, "A6")
    path = DOCS / "03_云帆仓配中心能效台账.xlsx"; wb.save(path); return path


def make_pdf(scene, chart):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = DOCS / "04_赤霄仓库应急响应手册.pdf"
    styles = getSampleStyleSheet()
    title = ParagraphStyle("CNTitle", parent=styles["Title"], fontName="STSong-Light", fontSize=24, leading=32, textColor=colors.HexColor("#17324d"), alignment=TA_CENTER)
    h1 = ParagraphStyle("CNH1", parent=styles["Heading1"], fontName="STSong-Light", fontSize=17, leading=24, textColor=colors.HexColor("#2f6690"))
    body = ParagraphStyle("CNBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=11, leading=18)
    cap = ParagraphStyle("CNCap", parent=body, alignment=TA_CENTER, textColor=colors.HexColor("#44545c"))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm,
                            title="赤霄仓库应急响应手册", author="RAG Benchmark")
    story = [Paragraph("赤霄仓库应急响应手册", title), Spacer(1, 6*mm),
             Paragraph("文档编号：NEBULA-PDF-2025-04　版本：2.1", cap), Spacer(1, 10*mm),
             Paragraph("1. 适用范围", h1),
             Paragraph("本手册适用于杭州临平赤霄仓库。总指挥为陆远，应急集合点位于东门银杏广场。发生全场断电时，值班员必须先广播口令“ORANGE-17”，再启动备用电源。", body),
             Spacer(1, 4*mm), Paragraph("2. 响应时限", h1),
             Paragraph("一级事件要求 5 分钟内确认、15 分钟内完成首轮处置；超过 30 分钟仍未恢复时必须通知区域负责人。", body),
             PageBreak(), Paragraph("3. 设备巡检责任表", h1)]
    rows = [["设备", "检查频率", "安全阈值", "责任人"], ["备用发电机", "每周一", "油量≥75%", "陆远"],
            ["冷链机组", "每日 08:30", "温度≤-18°C", "沈溪"], ["消防泵", "每月 5 日", "压力≥0.55MPa", "袁野"]]
    t = Table(rows, colWidths=[38*mm, 38*mm, 42*mm, 30*mm], repeatRows=1)
    t.setStyle(TableStyle([("FONTNAME", (0,0), (-1,-1), "STSong-Light"), ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2f6690")),
                           ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .5, colors.HexColor("#90a4ae")),
                           ("ALIGN", (0,0), (-1,-1), "CENTER"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#eef4f7")]),
                           ("FONTSIZE", (0,0), (-1,-1), 10), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
    story += [t, PageBreak(), Paragraph("4. 现场图与季度演练", h1), RLImage(str(scene), width=170*mm, height=97.1*mm), Paragraph("图 1　仓库巡检现场（需从图片识别对象）", cap), Spacer(1, 4*mm),
              RLImage(str(chart), width=170*mm, height=99.2*mm), Paragraph("图 2　季度演练参与人数", cap)]
    doc.build(story); return path


def write_ground_truth():
    items = [
        ("Q001", "docx", "title", "01_星云智能仓库运营报告.docx", "封面标题", "这份 Word 文档的完整标题是什么？", "星云智能仓库 2025 年运营报告", ["星云智能仓库", "2025", "运营报告"]),
        ("Q002", "docx", "body", "01_星云智能仓库运营报告.docx", "第1节 项目概况", "“苍穹”试点位于哪里，何时正式上线？", "苏州昆山；2025 年 3 月 18 日", ["苏州昆山", "2025年3月18日"]),
        ("Q003", "docx", "image", "01_星云智能仓库运营报告.docx", "图1", "巡检现场图中有几台蓝色机器人、几个橙色货箱？", "3 台蓝色机器人，2 个橙色货箱", ["3", "2", "蓝色", "橙色"]),
        ("Q004", "docx", "table", "01_星云智能仓库运营报告.docx", "第3节表格", "设备 C-03 的类型、数量、可用率和负责人分别是什么？", "充电桩，6 台，97.8%，韩梅", ["充电桩", "6", "97.8%", "韩梅"]),
        ("Q005", "docx", "chart", "01_星云智能仓库运营报告.docx", "图2", "季度订单量最高的是哪个季度？比第一季度多多少单？", "第四季度；多 600 单", ["第四季度", "600"]),
        ("Q006", "docx", "cross_modal", "01_星云智能仓库运营报告.docx", "图1 + 第3节表格", "图中蓝色机器人数量与表格中充电桩数量的最简比是多少？", "1:2", ["1:2"]),
        ("Q007", "pptx", "title", "02_北辰物流中心扩容方案.pptx", "幻灯片1标题", "PPT 的主标题是什么？", "北辰物流中心扩容方案", ["北辰物流中心扩容方案"]),
        ("Q008", "pptx", "body", "02_北辰物流中心扩容方案.pptx", "幻灯片1正文", "“极光”扩容计划在哪个作业区、何时启用？", "B3 区；2025 年 9 月 8 日", ["B3", "2025年9月8日"]),
        ("Q009", "pptx", "table", "02_北辰物流中心扩容方案.pptx", "幻灯片2表格", "试运行阶段的设备数、责任人和状态是什么？", "16 台，陈曦，进行中", ["16", "陈曦", "进行中"]),
        ("Q010", "pptx", "chart", "02_北辰物流中心扩容方案.pptx", "幻灯片3图表", "处理能力最高的是哪条线？数值是多少？", "B 线，510 箱/小时", ["B线", "510"]),
        ("Q011", "pptx", "body", "02_北辰物流中心扩容方案.pptx", "幻灯片3正文", "异常联系人的姓名和分机号是什么？", "陈曦，分机 6027", ["陈曦", "6027"]),
        ("Q012", "xlsx", "title", "03_云帆仓配中心能效台账.xlsx", "项目概览!A1:F2", "Excel 工作簿首页标题是什么？", "云帆仓配中心 2025 能效台账", ["云帆仓配中心", "2025", "能效台账"]),
        ("Q013", "xlsx", "body", "03_云帆仓配中心能效台账.xlsx", "项目概览!A4:B10", "夜间低谷时段是几点到几点？", "22:00–06:00", ["22:00", "06:00"]),
        ("Q014", "xlsx", "table", "03_云帆仓配中心能效台账.xlsx", "月度能耗!A1:E5", "3 月的订单量、耗电量和单位能耗分别是多少？", "1500 单，10200 kWh，6.80 kWh/单", ["1500", "10200", "6.80"]),
        ("Q015", "xlsx", "chart", "03_云帆仓配中心能效台账.xlsx", "月度能耗!G2", "单位能耗从 1 月到 4 月总体如何变化？下降了多少？", "持续下降；从 8.20 降至 6.20，共下降 2.00 kWh/单", ["下降", "8.20", "6.20", "2.00"]),
        ("Q016", "xlsx", "image", "03_云帆仓配中心能效台账.xlsx", "现场图片!A6", "现场图片中工人的安全帽是什么颜色，手里拿着什么？", "黄色安全帽；平板电脑", ["黄色", "平板"]),
        ("Q017", "pdf", "title", "04_赤霄仓库应急响应手册.pdf", "第1页标题", "PDF 的完整标题是什么？", "赤霄仓库应急响应手册", ["赤霄仓库", "应急响应手册"]),
        ("Q018", "pdf", "body", "04_赤霄仓库应急响应手册.pdf", "第1页第1节", "全场断电时先广播什么口令？集合点在哪里？", "口令 ORANGE-17；东门银杏广场", ["ORANGE-17", "东门银杏广场"]),
        ("Q019", "pdf", "table", "04_赤霄仓库应急响应手册.pdf", "第2页表格", "冷链机组的检查时间、安全阈值和责任人是什么？", "每日 08:30；温度≤-18°C；沈溪", ["08:30", "-18", "沈溪"]),
        ("Q020", "pdf", "chart", "04_赤霄仓库应急响应手册.pdf", "第3页图2", "季度演练参与人数最高的是哪一季度，共多少人？", "第四季度，96 人", ["第四季度", "96"]),
        ("Q021", "pdf", "image", "04_赤霄仓库应急响应手册.pdf", "第3页图1", "图片左侧充电区上方是什么颜色的标志，图案是什么？", "绿色标志，白色叶片图案", ["绿色", "白色", "叶"]),
        ("Q022", "multi", "cross_document", "全部文档", "DOCX第1节 + PPTX第1页 + XLSX项目概览 + PDF第1页", "四个文件中的负责人分别是谁？按 DOCX、PPTX、XLSX、PDF 顺序回答。", "林澈、陈曦、唐薇、陆远", ["林澈", "陈曦", "唐薇", "陆远"]),
    ]
    with (OUT / "ground_truth.jsonl").open("w", encoding="utf-8") as f:
        for qid, fmt, modality, source, locator, question, answer, keywords in items:
            f.write(json.dumps({"id": qid, "format": fmt, "modality": modality, "source": source, "locator": locator, "question": question, "answer": answer, "required_keywords": keywords}, ensure_ascii=False) + "\n")


def write_readme():
    text = """# 多模态 RAG 输入测试集\n\n本测试集用于验证 RAG 管线对 DOCX、PPTX、XLSX、PDF 的解析、切分、检索和回答能力，覆盖标题、正文、图片、表格、图表、跨模态及跨文件问题。\n\n## 文件\n\n- `documents/01_星云智能仓库运营报告.docx`：标题、分级正文、场景图片、原生表格、柱状图图片。\n- `documents/02_北辰物流中心扩容方案.pptx`：幻灯片标题/正文、场景图片、原生表格、原生图表。\n- `documents/03_云帆仓配中心能效台账.xlsx`：合并标题、正文单元格、数据表、原生图表、嵌入图片。\n- `documents/04_赤霄仓库应急响应手册.pdf`：分页标题/正文、表格、图片、图表。\n- `ground_truth.jsonl`：22 道问题及标准答案、定位信息、必含关键词。\n- `manifest.json`：文件清单与 SHA-256，便于检查传输完整性。\n- `assets/`：文档内嵌的确定性图片源文件。\n\n## 建议评测\n\n1. 将 `documents/` 作为唯一知识库输入，禁止把 `ground_truth.jsonl` 加入索引。\n2. 逐行读取 `ground_truth.jsonl` 的 `question` 发起查询。\n3. 检索层记录命中文档、页/幻灯片/工作表和 chunk；回答层检查 `required_keywords`。\n4. 建议报告：文档命中率、定位命中率、关键词全包含率、答案正确率、无依据回答率。\n5. 图片题必须开启 OCR/视觉模型；仅提取文本的管线预计无法通过 Q003、Q016、Q021。\n\n## 模态分布\n\n- 标题：4 题\n- 正文：5 题\n- 图片：3 题\n- 表格：4 题\n- 图表：4 题\n- 跨模态/跨文件：2 题\n\n## 评分口径\n\n精确字符串不是唯一标准。数字允许常见格式差异（如 `2025-03-18` 与 `2025 年 3 月 18 日`）；每题所有 `required_keywords` 均被答案语义覆盖可记为通过。\n\n## 图片生成说明\n\n在线图像生成服务两次网络失败，因此场景图改为 Pillow 本地确定性绘制。它仍包含可验证的对象、颜色、数量和人物动作，且四类文档共用同一图片，适合做解析一致性测试。\n"""
    (OUT / "README.md").write_text(text, encoding="utf-8")


def manifest(paths):
    import hashlib
    data = {"dataset": "multimodal-rag-office-v1", "version": "1.0", "files": []}
    for p in paths:
        b = p.read_bytes(); data["files"].append({"path": str(p.relative_to(OUT)).replace("\\", "/"), "bytes": len(b), "sha256": hashlib.sha256(b).hexdigest()})
    (OUT / "manifest.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    if OUT.exists(): shutil.rmtree(OUT)
    ASSETS.mkdir(parents=True); DOCS.mkdir(parents=True)
    scene = ASSETS / "warehouse_scene.png"
    orders = ASSETS / "quarterly_orders.png"
    drills = ASSETS / "quarterly_drill_participants.png"
    make_scene(scene)
    make_bar_chart(orders, "2025 年季度完成订单量", ["第一季度", "第二季度", "第三季度", "第四季度"], [1200, 1500, 1350, 1800])
    make_bar_chart(drills, "季度演练参与人数", ["第一季度", "第二季度", "第三季度", "第四季度"], [48, 72, 60, 96], color="#e07a5f")
    paths = [make_docx(scene, orders), make_pptx(scene), make_xlsx(scene), make_pdf(scene, drills)]
    write_ground_truth(); write_readme()
    paths += [OUT / "ground_truth.jsonl", OUT / "README.md", scene, orders, drills]
    manifest(paths)
    print(json.dumps({"output": str(OUT), "documents": [p.name for p in paths[:4]], "questions": 22}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
