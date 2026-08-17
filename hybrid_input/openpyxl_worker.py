from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

PACKAGE_ROOT = Path(__file__).resolve().parent.parent
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

from hybrid_input.contracts import Artifact, ArtifactBundle, Provenance


def _document_id(source: Path) -> str:
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _display_value(value: Any) -> str:
    value = _json_value(value)
    return "" if value is None else str(value)


def _markdown_value(value: Any) -> str:
    return _display_value(value).replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def _row_bands(worksheet: Any) -> list[tuple[int, int, int, int]]:
    populated: dict[int, list[int]] = {}
    for row in worksheet.iter_rows():
        columns = [cell.column for cell in row if cell.value is not None]
        if columns:
            populated[row[0].row] = columns
    if not populated:
        return []
    bands: list[tuple[int, int, int, int]] = []
    start = previous = min(populated)
    for row_number in sorted(populated)[1:]:
        if row_number > previous + 1:
            columns = [column for row in range(start, previous + 1) for column in populated.get(row, [])]
            bands.append((start, previous, min(columns), max(columns)))
            start = row_number
        previous = row_number
    columns = [column for row in range(start, previous + 1) for column in populated.get(row, [])]
    bands.append((start, previous, min(columns), max(columns)))
    return bands


def _merged_map(worksheet: Any) -> dict[tuple[int, int], Any]:
    result: dict[tuple[int, int], Any] = {}
    for merged in worksheet.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                result[(row, column)] = merged
    return result


def _region_provenance(
    source: Path,
    sheet: str,
    sheet_index: int,
    cell_range: str,
    bounds: tuple[int, int, int, int],
    max_row: int,
    max_column: int,
) -> Provenance:
    min_row, max_region_row, min_column, max_region_column = bounds
    return Provenance(
        source_file=source.name,
        page=sheet_index,
        sheet=sheet,
        cell_range=cell_range,
        bbox=[
            (min_column - 1) / max(max_column, 1),
            (min_row - 1) / max(max_row, 1),
            max_region_column / max(max_column, 1),
            max_region_row / max(max_row, 1),
        ],
        bbox_original=[float(min_column - 1), float(min_row - 1), float(max_region_column), float(max_region_row)],
        locator=f"sheet:{sheet}/range:{cell_range}",
    )


def _cell_records(
    worksheet: Any,
    bounds: tuple[int, int, int, int],
    merged_by_cell: dict[tuple[int, int], Any],
) -> list[dict[str, Any]]:
    min_row, max_row, min_column, max_column = bounds
    records: list[dict[str, Any]] = []
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            cell = worksheet.cell(row, column)
            merged = merged_by_cell.get((row, column))
            if merged and (row != merged.min_row or column != merged.min_col):
                continue
            if cell.value is None and not merged:
                continue
            record: dict[str, Any] = {
                "cell": cell.coordinate,
                "row": cell.row,
                "column": cell.column,
                "value": _json_value(cell.value),
                "row_span": merged.max_row - merged.min_row + 1 if merged else 1,
                "column_span": merged.max_col - merged.min_col + 1 if merged else 1,
            }
            if merged:
                record["merged_range"] = str(merged)
            if cell.data_type == "f":
                record["formula"] = str(cell.value)
            records.append(record)
    return records


def _table_markdown(worksheet: Any, bounds: tuple[int, int, int, int]) -> str:
    min_row, max_row, min_column, max_column = bounds
    rows = [
        [_markdown_value(worksheet.cell(row, column).value) for column in range(min_column, max_column + 1)]
        for row in range(min_row, max_row + 1)
    ]
    if not rows:
        return ""
    lines = ["| " + " | ".join(row) + " |" for row in rows]
    lines.insert(1, "| " + " | ".join("---" for _ in rows[0]) + " |")
    return "\n".join(lines)


def _is_text_region(records: list[dict[str, Any]], bounds: tuple[int, int, int, int]) -> bool:
    min_row, max_row, _, _ = bounds
    counts = {
        row: sum(1 for record in records if record["row"] == row)
        for row in range(min_row, max_row + 1)
    }
    return bool(records) and all(count <= 1 for count in counts.values())


def _anchor_bounds(anchor: Any) -> tuple[int, int, int, int]:
    if isinstance(anchor, str):
        from openpyxl.utils.cell import coordinate_to_tuple

        row, column = coordinate_to_tuple(anchor)
        return row, row, column, column
    start = getattr(anchor, "_from", None)
    end = getattr(anchor, "to", None)
    if start is None:
        return 1, 1, 1, 1
    min_row, min_column = int(start.row) + 1, int(start.col) + 1
    if end is None:
        return min_row, min_row, min_column, min_column
    return min_row, int(end.row) + 1, min_column, int(end.col) + 1


def convert(source: Path, output_dir: Path) -> ArtifactBundle:
    import openpyxl
    from openpyxl.utils import get_column_letter

    source = source.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    raw_image_dir = output_dir / "raw-images"
    raw_image_dir.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False)
    doc_id = _document_id(source)
    staged: list[tuple[int, int, int, int, Artifact]] = []
    warnings: list[str] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            image_entries = [(image, _anchor_bounds(image.anchor)) for image in getattr(worksheet, "_images", [])]
            chart_entries = [(chart, _anchor_bounds(chart.anchor)) for chart in getattr(worksheet, "_charts", [])]
            object_bounds = [bounds for _, bounds in image_entries + chart_entries]
            max_row = max([worksheet.max_row, 1] + [bounds[1] for bounds in object_bounds])
            max_column = max([worksheet.max_column, 1] + [bounds[3] for bounds in object_bounds])
            merged_by_cell = _merged_map(worksheet)
            for min_row, max_region_row, min_column, max_region_column in _row_bands(worksheet):
                bounds = (min_row, max_region_row, min_column, max_region_column)
                cell_range = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_region_column)}{max_region_row}"
                provenance = _region_provenance(
                    source, worksheet.title, sheet_index, cell_range, bounds, max_row, max_column
                )
                records = _cell_records(worksheet, bounds, merged_by_cell)
                merged_ranges = sorted(
                    {
                        record["merged_range"]
                        for record in records
                        if "merged_range" in record
                    }
                )
                metadata = {
                    "semantic_type": "worksheet_region",
                    "parser": "openpyxl",
                    "structure_origin": "native",
                    "span_reliability": "exact",
                    "cell_range": cell_range,
                    "cells": records,
                    "merged_ranges": merged_ranges,
                    "rows": max_region_row - min_row + 1,
                    "columns": max_region_column - min_column + 1,
                }
                if _is_text_region(records, bounds):
                    for row_order, record in enumerate(records):
                        cell = worksheet[record["cell"]]
                        single_bounds = (cell.row, cell.row, cell.column, cell.column + int(record["column_span"]) - 1)
                        single_range = record.get("merged_range") or record["cell"]
                        single_provenance = _region_provenance(
                            source, worksheet.title, sheet_index, str(single_range), single_bounds, max_row, max_column
                        )
                        staged.append(
                            (
                                sheet_index,
                                cell.row,
                                cell.column,
                                row_order,
                                Artifact(
                                    "pending",
                                    "text",
                                    single_provenance,
                                    text=_display_value(record["value"]),
                                    metadata={
                                        "semantic_type": "worksheet_text",
                                        "parser": "openpyxl",
                                        "cell": record,
                                    },
                                ),
                            )
                        )
                else:
                    staged.append(
                        (
                            sheet_index,
                            min_row,
                            min_column,
                            0,
                            Artifact(
                                "pending",
                                "table",
                                provenance,
                                text=_table_markdown(worksheet, bounds),
                                metadata=metadata,
                            ),
                        )
                    )

            for image_index, (image, bounds) in enumerate(image_entries, start=1):
                min_row, max_image_row, min_column, max_image_column = bounds
                cell_range = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_image_column)}{max_image_row}"
                extension = str(getattr(image, "format", None) or "png").lower()
                destination = raw_image_dir / f"sheet-{sheet_index:03d}-image-{image_index:03d}.{extension}"
                try:
                    destination.write_bytes(image._data())
                except Exception as exc:
                    warnings.append(f"Embedded Excel image extraction failed on {worksheet.title}: {exc}")
                    continue
                provenance = _region_provenance(
                    source, worksheet.title, sheet_index, cell_range, bounds, max_row, max_column
                )
                staged.append(
                    (
                        sheet_index,
                        min_row,
                        min_column,
                        1,
                        Artifact(
                            "pending",
                            "image",
                            provenance,
                            asset_path=str(destination.resolve()),
                            metadata={"semantic_type": "embedded_image", "parser": "openpyxl"},
                        ),
                    )
                )

            for chart_index, (chart, bounds) in enumerate(chart_entries, start=1):
                min_row, max_chart_row, min_column, max_chart_column = bounds
                cell_range = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_chart_column)}{max_chart_row}"
                provenance = _region_provenance(
                    source, worksheet.title, sheet_index, cell_range, bounds, max_row, max_column
                )
                staged.append(
                    (
                        sheet_index,
                        min_row,
                        min_column,
                        2,
                        Artifact(
                            "pending",
                            "visual_task",
                            provenance,
                            metadata={
                                "semantic_type": "excel_chart",
                                "parser": "openpyxl",
                                "chart_index": chart_index,
                                "render_required": True,
                            },
                        ),
                    )
                )
    finally:
        workbook.close()

    artifacts: list[Artifact] = []
    for order, (_, _, _, _, artifact) in enumerate(sorted(staged, key=lambda item: item[:4])):
        artifact.block_id = f"{doc_id[:16]}:{order:06d}"
        artifact.reading_order = order
        artifacts.append(artifact)
    if not artifacts:
        warnings.append("openpyxl parser produced no worksheet content")
    return ArtifactBundle(doc_id, str(source), "xlsx", "openpyxl-subprocess", artifacts, warnings)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--result", required=True, type=Path)
    args = parser.parse_args()
    bundle = convert(args.input, args.output)
    args.result.parent.mkdir(parents=True, exist_ok=True)
    args.result.write_text(json.dumps(bundle.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
